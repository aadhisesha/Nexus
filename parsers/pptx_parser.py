"""
parsers/pptx_parser.py — PowerPoint extraction via python-pptx.

Public API
----------
extract_text(filepath) -> str
    Returns slide-by-slide text (all text frames + speaker notes).
    On any error returns "" — never raises.
"""

from __future__ import annotations

import logging

log = logging.getLogger(__name__)


def extract_text(filepath: str) -> str:
    """
    Extract plain text from a .pptx file, slide by slide.

    Parameters
    ----------
    filepath : str
        Absolute or relative path to the .pptx file.

    Returns
    -------
    str
        All slide text concatenated, or "" on failure.
    """
    try:
        try:
            from pptx import Presentation
        except ImportError:
            log.error("python-pptx is not installed. Run: pip install python-pptx")
            return ""

        try:
            prs = Presentation(filepath)
        except Exception as exc:
            log.warning("Cannot open PPTX '%s' (corrupted or unsupported format): %s", filepath, exc)
            return ""

        slide_parts: list[str] = []
        total_slides = 0

        try:
            slides = list(prs.slides)
            total_slides = len(slides)
        except Exception as exc:
            log.warning("Error reading slides in PPTX '%s': %s", filepath, exc)
            slides = []

        for slide_num, slide in enumerate(slides, start=1):
            lines: list[str] = [f"=== Slide {slide_num} ==="]

            # --- Text frames in all shapes ---------------------------------------
            try:
                for shape in slide.shapes:
                    if not getattr(shape, "has_text_frame", False):
                        continue
                    for para in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in para.runs if hasattr(run, "text")).strip()
                        if not text and hasattr(para, "text"):
                            text = para.text.strip()
                        if text:
                            lines.append(text)
            except Exception as exc:
                log.warning("Error reading shapes on slide %d in PPTX '%s': %s", slide_num, filepath, exc)

            # --- Speaker notes ---------------------------------------------------
            try:
                if getattr(slide, "has_notes_slide", False) or hasattr(slide, "notes_slide"):
                    notes_slide = slide.notes_slide
                    if notes_slide and hasattr(notes_slide, "notes_text_frame"):
                        notes_tf = notes_slide.notes_text_frame
                        notes_text = notes_tf.text.strip() if notes_tf else ""
                        if notes_text:
                            lines.append("[Notes]")
                            lines.append(notes_text)
            except Exception:
                pass  # Notes not present — safe to ignore

            if len(lines) > 1:
                slide_parts.append("\n".join(lines))

        result = "\n\n".join(slide_parts).strip()
        if not result:
            log.warning("PPTX '%s' (%d slides) contains no text (empty or image-only slides).", filepath, total_slides)

        return result

    except Exception as exc:
        log.warning("Unhandled exception in pptx_parser for '%s': %s", filepath, exc)
        return ""
