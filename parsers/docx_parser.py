"""
parsers/docx_parser.py — Word document extraction via python-docx.

Public API
----------
extract_text(filepath) -> str
    Returns all paragraph text + table cell text joined with newlines.
    On any error returns "" — never raises.
"""

from __future__ import annotations

import logging

log = logging.getLogger(__name__)


def extract_text(filepath: str) -> str:
    """
    Extract plain text from a .docx file.

    Parameters
    ----------
    filepath : str
        Absolute or relative path to the .docx file.

    Returns
    -------
    str
        Combined paragraph + table text, or "" on failure.
    """
    try:
        try:
            from docx import Document
        except ImportError:
            log.error("python-docx is not installed. Run: pip install python-docx")
            return ""

        try:
            doc = Document(filepath)
        except Exception as exc:
            log.warning("Cannot open DOCX '%s' (corrupted or unsupported format): %s", filepath, exc)
            return ""

        parts: list[str] = []

        # --- Paragraphs ----------------------------------------------------------
        try:
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    parts.append(text)
        except Exception as exc:
            log.warning("Error reading paragraphs in DOCX '%s': %s", filepath, exc)

        # --- Tables --------------------------------------------------------------
        try:
            for table_idx, table in enumerate(doc.tables, start=1):
                table_lines: list[str] = [f"[Table {table_idx}]"]
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    # Merge adjacent duplicate cells (merged cells repeat in python-docx)
                    deduped: list[str] = []
                    for cell_text in row_cells:
                        if not deduped or deduped[-1] != cell_text:
                            deduped.append(cell_text)
                    if any(deduped):
                        table_lines.append("\t".join(deduped))
                if len(table_lines) > 1:
                    parts.append("\n".join(table_lines))
        except Exception as exc:
            log.warning("Error reading tables in DOCX '%s': %s", filepath, exc)

        result = "\n".join(parts).strip()
        if not result:
            log.warning("DOCX '%s' contains no extractable text (empty or image-only document).", filepath)

        return result

    except Exception as exc:
        log.warning("Unhandled exception in docx_parser for '%s': %s", filepath, exc)
        return ""
